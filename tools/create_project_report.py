from pathlib import Path
import json, hashlib, re
from collections import defaultdict
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports'
OUT.mkdir(exist_ok=True)
DOC=OUT/'Bao_cao_ma_nguon_va_giao_dien_GalleryApp.docx'
SHOTS=ROOT/'build/previews/report'
descriptions={
 'AccountActivity':'Đọc hồ sơ users/{uid} từ Firestore, cho phép sửa họ tên và giới thiệu; lưu bằng merge và đồng bộ tên Firebase Auth.',
 'AlbumActivity':'Màn danh sách album; kế thừa xử lý chung từ GalleryActivity.',
 'AlbumDetailActivity':'Màn nội dung album, nhận albumId và dùng xử lý thêm ảnh, đổi tên, chọn nhiều từ GalleryActivity.',
 'AuthActivity':'Xử lý đăng nhập và đăng ký email hoặc số điện thoại, xác minh email, gửi SMS, đặt lại mật khẩu và ghi hồ sơ Firestore.',
 'BackupActivity':'Giữ màn Sao lưu và đồng bộ để minh họa. Nút sao lưu giải thích trạng thái demo; không tải file lên Storage.',
 'BaseActivity':'Thiết lập giao diện chung, xử lý khoảng cách thanh hệ thống và bàn phím, điều hướng bốn tab, hộp thoại và repository.',
 'CreateAlbumActivity':'Màn tạo album; dùng GalleryActivity để chọn ảnh, nhập tên và ghi album cùng quan hệ thành viên.',
 'EditPhotoActivity':'Điều phối công cụ cắt, xoay, bộ lọc, sáng và tương phản, vẽ, hoàn tác và lưu bản sao; xác nhận khi bỏ thay đổi.',
 'GalleryActivity':'Xử lý thư viện, tìm kiếm, lọc, sắp xếp, nhập tệp, chọn nhiều theo thứ tự, ẩn và bỏ ẩn, album và thùng rác.',
 'GalleryApplication':'Khởi tạo ứng dụng, áp dụng chế độ sáng tối và chuẩn bị Firebase để màn chào có thể hiển thị sớm.',
 'GalleryRepository':'Quản lý cache riêng của khách hoặc tài khoản, file nội bộ và metadata Firestore. Có mã Storage dự phòng, nhưng màn sao lưu của bản demo không gọi upload.',
 'GoogleSignInActivity':'Mở luồng chọn tài khoản Google thật, lấy thông tin xác thực và đăng nhập Firebase.',
 'HiddenMediaActivity':'Màn riêng để xem, tìm và lọc ảnh hoặc video đã ẩn; kế thừa thao tác bỏ ẩn và chọn nhiều.',
 'LoginActivity':'Điểm vào màn đăng nhập; kế thừa AuthActivity.',
 'MainActivity':'Điểm vào thư viện Ảnh; kế thừa GalleryActivity.',
 'MediaAdapter':'Hiển thị ô ảnh, video, album và danh sách. Nạp hình bằng Glide, đếm thành viên album và hiển thị huy hiệu số theo thứ tự chọn.',
 'NotificationsActivity':'Lưu và đọc lịch sử hoạt động cục bộ và Firestore; hỗ trợ đánh dấu đã đọc. Không phải thông báo đẩy FCM.',
 'OtpActivity':'Xử lý các ô mã xác minh SMS và bước kiểm tra liên kết xác minh email; cho phép gửi lại theo luồng tương ứng.',
 'PhotoDetailActivity':'Hiển thị ảnh, vuốt giữa các mục cùng trạng thái ẩn, yêu thích, thông tin, sửa, chia sẻ, ẩn và thùng rác; mở video bằng trình phát hệ thống.',
 'RegisterActivity':'Điểm vào màn đăng ký; kế thừa AuthActivity.',
 'SearchActivity':'Màn tìm kiếm dùng chung xử lý lọc và sắp xếp từ GalleryActivity.',
 'SettingsActivity':'Màn Hồ sơ và Cài đặt; mở tài khoản, giao diện, dung lượng, thông báo, các mục đã ẩn, thùng rác và đăng xuất.',
 'ShareActivity':'Chia sẻ file bằng FileProvider và bảng chia sẻ Android. Liên kết chỉ dùng được khi mục đã có URL cloud hợp lệ.',
 'SplashActivity':'Màn chào; đưa người dùng vào đăng nhập hoặc thư viện tùy trạng thái tài khoản. Có lựa chọn khám phá trên thiết bị.',
 'ThemeManager':'Lưu lựa chọn sáng tối và áp dụng AppCompatDelegate cho toàn bộ ứng dụng.',
 'TrashActivity':'Màn thùng rác; kế thừa xử lý khôi phục, xóa vĩnh viễn và dọn mục quá hạn từ GalleryActivity.',
 'EditorCanvasView':'Vẽ ảnh xem trước, khung cắt và lưới một phần ba; nhận kéo góc, di chuyển khung và nét vẽ theo tọa độ ảnh.',
 'EditorDocument':'Lưu trạng thái ảnh và nét vẽ, lịch sử hoàn tác, biến đổi cắt xoay, bộ lọc và kết xuất bitmap.',
 'EditorViewModel':'Giữ phiên chỉnh sửa khi Activity được tạo lại; giải mã, biến đổi và xuất JPEG trên luồng nền; lưu thành bản sao.',
 'FirebaseHelper':'Lớp hỗ trợ Firebase từ phiên bản ban đầu, gồm thao tác ảnh, album và thông báo. Các Activity hiện tại chủ yếu dùng GalleryRepository, không gọi helper này.',
 'PermissionHelper':'Các hàm hỗ trợ kiểm tra và yêu cầu quyền của bản mã ban đầu. Luồng nhập hiện tại dùng trình chọn tài liệu Android, không yêu cầu đọc toàn bộ thư viện.',
 'PhotoHelper':'Các tiện ích cho ảnh và video, định dạng kích thước hoặc thời gian và xử lý metadata.',
 'Album':'Mô hình album: mã, chủ sở hữu, tên, ảnh bìa và mốc thời gian.',
 'Notification':'Mô hình dữ liệu thông báo hoạt động.',
 'Photo':'Mô hình ảnh hoặc video, gồm tệp, kích thước, loại, trạng thái yêu thích, ẩn, thùng rác và danh sách album.',
 'User':'Mô hình hồ sơ người dùng và các thuộc tính tài khoản.'
}
dir_roles={
 'com/example/appanh':'Các Activity, điều hướng, xác thực, repository, adapter và cấu hình dùng chung.',
 'com/example/appanh/editor':'Bộ xử lý chỉnh sửa ảnh, vùng tương tác và trạng thái giữ qua vòng đời Activity.',
 'com/example/appanh/helpers':'Các tiện ích và lớp hỗ trợ; có helper Firebase và quyền từ mã ban đầu.',
 'com/example/appanh/models':'Các lớp mô hình dữ liệu dùng cho hồ sơ, ảnh, album và thông báo.',
 'layout':'Bố cục màn hình, các sheet, hộp thoại và ô dữ liệu.',
 'drawable':'Vector icon, hình minh họa và nền bo góc hoặc khung chọn.',
 'drawable-night':'Tài nguyên drawable thay thế khi giao diện tối được bật.',
 'color':'Màu có trạng thái cho chip và thanh điều hướng.',
 'menu':'Khai báo menu thanh điều hướng và hành động trên toolbar.',
 'values':'Chuỗi, màu, theme và style mặc định của ứng dụng.',
 'values-night':'Màu, theme và giá trị thay thế trong chế độ tối.',
 'values-v31':'Thiết lập theme dành cho Android API 31 trở lên.',
 'xml':'Chính sách sao lưu hệ thống và các đường dẫn chia sẻ FileProvider.',
 'mipmap-anydpi-v26':'Cấu hình biểu tượng ứng dụng thích ứng từ Android API 26.'
}
layout_roles={
 'activity_splash':'Màn chào có logo, minh họa, nút Bắt đầu và khám phá trên thiết bị.',
 'activity_login':'Biểu mẫu đăng nhập, quên mật khẩu, đăng ký và tiếp tục với Google.',
 'activity_register':'Biểu mẫu tạo tài khoản và chuyển về đăng nhập.',
 'activity_otp':'Màn xác thực với các ô nhập mã và thông tin gửi lại.',
 'activity_main':'Trang Ảnh với toolbar, tìm kiếm, chip lọc, lưới ảnh, nút nhập và thanh điều hướng.',
 'activity_search':'Trang tìm kiếm và lọc trong thư viện.',
 'activity_album':'Danh sách album hai cột với nút tạo album.',
 'activity_album_detail':'Nội dung một album và các hành động trên toolbar.',
 'activity_create_album':'Nhập tên và chọn ảnh cho album mới.',
 'activity_photo_detail':'Màn xem ảnh hoặc video, yêu thích, chỉnh sửa, chia sẻ và xóa.',
 'activity_edit_photo':'Bố trí vùng chỉnh sửa tùy biến, công cụ, bảng màu, tỷ lệ cắt và nút lưu bản sao.',
 'activity_account':'Hồ sơ tài khoản đọc và ghi Firestore.',
 'activity_settings':'Trang Hồ sơ và Cài đặt, các dòng chức năng có mô tả phụ.',
 'activity_backup':'Giao diện minh họa sao lưu với hai công tắc và nút sao lưu.',
 'activity_notifications':'Danh sách thông báo hoạt động và nút đánh dấu đã đọc.',
 'activity_trash':'Danh sách mục trong thùng rác và nút dọn.',
 'activity_google_sign_in':'Trang khởi động luồng đăng nhập Google.',
 'activity_dark_gallery':'Bố cục tham chiếu giao diện tối; trang Ảnh thực tế dùng activity_main với tài nguyên theo theme.',
 'item_media':'Ô ảnh hoặc video trong lưới; huy hiệu số thể hiện thứ tự chọn.',
 'item_album':'Thẻ album có ảnh bìa, tên, số lượng và nút ba chấm.',
 'item_media_row':'Một mục dạng danh sách, dùng trong thùng rác.',
 'item_notification':'Một thông báo hoạt động gồm biểu tượng và nội dung.',
 'dialog_photo_info':'Thông tin tên, kích thước, thời gian, vị trí lưu và loại ảnh.',
 'sheet_options':'Menu chọn nhiều, thêm album, ẩn, xem các mục đã ẩn và cài đặt.',
 'sheet_photo_context':'Menu khi nhấn giữ ảnh hoặc video.',
 'sheet_album_options':'Menu chọn nhiều, thêm ảnh, đổi tên, bỏ thành viên hoặc xóa album.',
 'sheet_filter':'Các cách sắp xếp thư viện.',
 'sheet_share':'Các lựa chọn chia sẻ tệp hoặc liên kết nếu đã có URL hợp lệ.'
}

doc=Document()
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.8)
for style in doc.styles:
    for border in list(style.element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
    if style.type in (1,2,3):
        style.font.name='Times New Roman';style.font.size=Pt(13);style.font.color.rgb=RGBColor(0,0,0)
        rp=style.element.get_or_add_rPr();fonts=rp.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rp.insert(0,fonts)
        for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+attr),None)
        for color in rp.findall(qn('w:color')):
            color.attrib.clear();color.set(qn('w:val'),'000000')
normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.size=Pt(13);st.font.bold=name!='Subtitle';st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
code=doc.styles.add_style('SourceCode',1);code.base_style=normal;code.font.name='Times New Roman';code.font.size=Pt(13);code.font.color.rgb=RGBColor(0,0,0)
code.paragraph_format.line_spacing=1;code.paragraph_format.space_before=Pt(0);code.paragraph_format.space_after=Pt(0);code.paragraph_format.keep_together=False;code.paragraph_format.widow_control=False
pp=code.element.get_or_add_pPr();wrap=OxmlElement('w:wordWrap');wrap.set(qn('w:val'),'0');pp.append(wrap)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Trang ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Mã nguồn và giao diện ứng dụng Ảnh GalleryApp'
doc.core_properties.subject='Logic Java giao diện XML và kết quả chạy ứng dụng Android'
doc.core_properties.author=''

def para(text='',bold=False):
    p=doc.add_paragraph();p.add_run(text).bold=bold;return p
def heading(text,level=1):return doc.add_heading(text,level)
manifest=[]
def source_file(path,index,purpose):
    text=path.read_text(encoding='utf-8-sig');text=text.replace('\r\n','\n')
    heading(index+' '+path.stem,3)
    para('Nguồn: '+path.relative_to(ROOT).as_posix())
    para('Vai trò: '+purpose)
    para('Mã nguồn đầy đủ',True)
    for line in text.split('\n'):doc.add_paragraph(line,style='SourceCode')
    para()
    manifest.append({'path':path.relative_to(ROOT).as_posix(),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'characters':len(text),'lines':len(text.splitlines())})

doc.add_paragraph('Mã nguồn và giao diện ứng dụng Ảnh GalleryApp',style='Title')
para('Báo cáo logic xử lý giao diện và kết quả chạy ứng dụng Android',True)
para('Ứng dụng Ảnh quản lý ảnh và video, tạo album, tìm kiếm, ẩn hoặc bỏ ẩn, chỉnh sửa ảnh và quản lý hồ sơ tài khoản. Báo cáo trình bày ba phần: mã Java xử lý, mã XML giao diện và ảnh chụp kết quả chạy trên máy ảo Gallery Light Android 11 API 30.')
para('Mã nguồn được lấy từ project GalleryApp tại thời điểm lập báo cáo ngày 18 tháng 10 năm 2026. Phần Java gồm toàn bộ 36 file trong app/src/main/java. Phần XML gồm AndroidManifest.xml và toàn bộ 154 file XML trong app/src/main/res; không lấy file tự sinh trong build hoặc mã kiểm thử.')
para('Bản hiện tại lưu ảnh và video trong bộ nhớ riêng của ứng dụng. Hồ sơ và metadata album sử dụng Firestore. Màn Sao lưu và đồng bộ giữ giao diện minh họa, chưa thực hiện sao lưu cloud. Chức năng ẩn là bộ lọc trong ứng dụng, chưa có khóa mật khẩu.')
para('Nội dung',True)
for t in ['1 Logic xử lý','2 Giao diện ứng dụng','3 Kết quả sau khi chạy']:para(t)

doc.add_page_break();heading('1 Logic xử lý')
para('Thư mục gốc chứa mã Java: app/src/main/java/com/example/appanh. Các Activity xử lý vòng đời màn hình và sự kiện người dùng; GalleryRepository quản lý dữ liệu; các thư mục editor, models và helpers tách riêng công cụ chỉnh sửa, mô hình và tiện ích.')
para('Luồng thư viện: Activity nhận thao tác → GalleryActivity cập nhật trạng thái tìm kiếm hoặc chọn → MediaAdapter hiển thị dữ liệu → GalleryRepository lưu cache và metadata khi phù hợp. Chọn nhiều dùng LinkedHashSet để giữ thứ tự; adapter hiển thị số 1, 2, 3 theo thứ tự đó.')
para('Luồng chỉnh sửa: EditPhotoActivity nhận photoId → EditorViewModel đọc ảnh trên luồng nền → EditorCanvasView nhận thao tác cắt hoặc vẽ → EditorDocument tạo trạng thái mới → xuất JPEG thành tệp khác và thêm vào thư viện, giữ nguyên ảnh gốc.')
java=defaultdict(list)
for path in sorted((ROOT/'app/src/main/java').rglob('*.java')):java[path.parent.relative_to(ROOT/'app/src/main/java').as_posix()].append(path)
for i,(folder,files) in enumerate(java.items(),1):
    heading(f'1 {i} Thư mục '+folder.split('/')[-1],2)
    para('Đường dẫn: app/src/main/java/'+folder)
    para(dir_roles.get(folder,'Mã nguồn xử lý của ứng dụng.')+f' Số file Java: {len(files)}.')
    para('Danh sách: '+', '.join(p.name for p in files)+'.')
    for j,path in enumerate(files,1):source_file(path,f'1 {i} {j}',descriptions[path.stem])

doc.add_page_break();heading('2 Giao diện ứng dụng')
para('Nguồn giao diện nằm trong app/src/main/res. Layout XML mô tả vị trí và cấu trúc thành phần; menu khai báo hành động; drawable chứa icon và nền; color và values cung cấp màu và style. Các thư mục night chứa tài nguyên thay thế để cùng bố cục chuyển sang chế độ tối.')
para('HiddenMediaActivity sử dụng lại activity_main.xml và thay tiêu đề cùng trạng thái lọc; không có một layout XML riêng. Vùng ảnh của activity_edit_photo.xml sử dụng EditorCanvasView viết bằng Java để hiển thị khung cắt và nét vẽ.')
heading('2 1 Khai báo ứng dụng',2)
source_file(ROOT/'app/src/main/AndroidManifest.xml','2 1 1','Khai báo Application, các Activity, quyền mạng, FileProvider và chính sách chia sẻ hoặc sao lưu dữ liệu hệ thống.')
xml=defaultdict(list)
for path in sorted((ROOT/'app/src/main/res').rglob('*.xml')):xml[path.parent.relative_to(ROOT/'app/src/main/res').as_posix()].append(path)
order=['layout','menu','values','values-night','values-v31','color','drawable','drawable-night','mipmap-anydpi-v26','xml']
for i,folder in enumerate(sorted(xml,key=lambda k:(order.index(k) if k in order else 99,k)),2):
    files=xml[folder];role=dir_roles.get(folder,'Tài nguyên XML dùng trong ứng dụng.')
    heading(f'2 {i} Thư mục '+folder.replace('-',' '),2)
    para('Đường dẫn: app/src/main/res/'+folder);para(role+f' Số file XML: {len(files)}.')
    para('Danh sách: '+', '.join(p.name for p in files)+'.')
    for j,path in enumerate(files,1):
        purpose=layout_roles.get(path.stem,role+' Tên tài nguyên: '+path.stem+'.')
        source_file(path,f'2 {i} {j}',purpose)

doc.add_page_break();heading('3 Kết quả sau khi chạy')
para('Các hình dưới đây là ảnh chụp ứng dụng chạy trên máy ảo, không phải bản thiết kế tham khảo. Mỗi hình đi kèm tên bố cục, lớp xử lý, thao tác và ý nghĩa để trình bày các giao diện chính.')
para('Số lượng ảnh hoặc album phản ánh dữ liệu demo có trên máy tại lúc chụp. Một số màn hình có dữ liệu tài khoản đang đăng nhập; tên và số mục không được gán cố định trong giao diện.')
screens=[
 ('Màn hình khởi động','SplashActivity.png','activity_splash.xml','SplashActivity','Hiển thị logo, tên ứng dụng, minh họa và hai lựa chọn truy cập.','Nhấn Bắt đầu để đi tiếp theo trạng thái đăng nhập; Khám phá trên thiết bị mở chế độ khách. Không cần thao tác đăng nhập để giới thiệu bố cục màn chào.'),
 ('Màn hình đăng nhập','LoginActivity.png','activity_login.xml','LoginActivity và AuthActivity','Biểu mẫu email hoặc số điện thoại, mật khẩu, quên mật khẩu và tiếp tục với Google.','Nhập tài khoản rồi nhấn Đăng nhập. Email và số điện thoại đi theo hai cách xác thực khác nhau; cửa sổ Google là luồng thật của hệ thống.'),
 ('Màn hình đăng ký','RegisterActivity.png','activity_register.xml','RegisterActivity và AuthActivity','Biểu mẫu họ tên, email hoặc số điện thoại và mật khẩu; có đường dẫn quay về đăng nhập.','Các trường được kiểm tra trước khi gọi Firebase. Đăng ký email gửi liên kết xác minh; không dùng mã email giả trong bản demo.'),
 ('Thư viện và thứ tự chọn nhiều','MainActivity.png','activity_main.xml và item_media.xml','MainActivity GalleryActivity và MediaAdapter','Thanh tìm kiếm, bộ lọc, lưới ảnh, nút thêm tệp và bốn tab điều hướng. Huy hiệu 1, 2, 3 thể hiện thứ tự chạm chọn, không phải vị trí trong lưới.','Mở menu ba chấm → Chọn nhiều → chạm các mục. Bỏ một mục thì các số phía sau cập nhật; chọn lại thì mục đó được xếp cuối. Có thể thêm các mục đã chọn vào album hoặc ẩn.'),
 ('Danh sách album','AlbumActivity.png','activity_album.xml và item_album.xml','AlbumActivity GalleryActivity và MediaAdapter','Các thẻ album hai cột gồm ảnh bìa thật, tên và số lượng ảnh hoặc video.','Nhấn Tạo album để nhóm ảnh. Nhấn thẻ để xem nội dung; menu ba chấm hỗ trợ thêm ảnh, đổi tên hoặc xóa album. Xóa album không xóa tệp trong thư viện.'),
 ('Nội dung album Du lịch','AlbumDetailActivity.png','activity_album_detail.xml','AlbumDetailActivity và GalleryActivity','Lưới các ảnh thuộc album cùng nút thêm và menu hành động trên toolbar.','Nhấn + để chọn thêm ảnh từ thư viện. Menu hỗ trợ chọn nhiều và bỏ các mục khỏi album; quan hệ thành viên được lưu trong albumIds.'),
 ('Tạo album mới','CreateAlbumActivity.png','activity_create_album.xml','CreateAlbumActivity và GalleryActivity','Ô tên album và lưới chọn ảnh; các ảnh được chọn có số thứ tự.','Nhập tên không rỗng, chọn ảnh rồi nhấn tạo. Album mới được ghi vào thư viện của khách hoặc tài khoản đang dùng.'),
 ('Tìm kiếm ảnh và video','SearchActivity.png','activity_search.xml','SearchActivity và GalleryActivity','Ô tìm theo tên, chip loại nội dung và bộ lọc sắp xếp.','Gõ một phần tên tệp; chọn Ảnh hoặc Video để thu hẹp kết quả. Mục đã ẩn không xuất hiện trong kết quả tìm kiếm thường.'),
 ('Xem ảnh chi tiết','PhotoDetailActivity.png','activity_photo_detail.xml','PhotoDetailActivity','Ảnh được hiển thị lớn với hành động yêu thích, chỉnh sửa, chia sẻ và thùng rác.','Vuốt để chuyển ảnh; menu bổ sung mở thông tin hoặc ẩn và bỏ ẩn. Khi mở video, chạm vùng xem trước để phát bằng ứng dụng hệ thống.'),
 ('Cắt ảnh bằng khung tương tác','editor-crop.png','activity_edit_photo.xml','EditPhotoActivity EditorCanvasView và EditorDocument','Khung cắt có bốn góc kéo, lưới một phần ba và vùng ngoài được làm tối; có tỷ lệ tự do, 1:1, 4:3 và 16:9.','Chọn Cắt → kéo góc để thay kích thước hoặc kéo bên trong để di chuyển → Áp dụng cắt. Có thể xoay 90 độ; Lưu bản sao tạo tệp mới thay vì ghi đè ảnh gốc.'),
 ('Vẽ trực tiếp lên ảnh','editor-draw.png','activity_edit_photo.xml','EditPhotoActivity EditorCanvasView và EditorDocument','Bảng sáu màu, thanh độ dày và nút xóa nét. Nét vẽ hiển thị trên ảnh theo tọa độ gốc.','Chọn Vẽ → chọn màu và độ dày → kéo ngón tay hoặc chuột trên ảnh. Hoàn tác và Làm lại áp dụng cho thao tác chỉnh sửa; bản xuất JPEG bao gồm cả ảnh và nét vẽ.'),
 ('Ảnh và video đã ẩn','HiddenMediaActivity.png','Dùng lại activity_main.xml và item_media.xml','HiddenMediaActivity và GalleryActivity','Màn riêng có tiêu đề, tìm kiếm, lọc ảnh hoặc video và danh sách các mục đã ẩn.','Mở Hồ sơ → Ảnh và video đã ẩn hoặc menu trang Ảnh. Nhấn giữ một mục → Bỏ ẩn để đưa về thư viện; có chọn nhiều. Chức năng này chưa có khóa mật khẩu.'),
 ('Hồ sơ và cài đặt','SettingsActivity.png','activity_settings.xml','SettingsActivity','Các dòng chức năng gồm tài khoản, sao lưu, dung lượng, thông báo, giao diện, ngôn ngữ và nhóm thư viện phía dưới.','Nhấn từng dòng để mở màn tương ứng. Chọn Giao diện để đổi sáng hoặc tối cho toàn app; lựa chọn được giữ khi mở lại ứng dụng.'),
 ('Thông tin tài khoản trên Firestore','AccountActivity.png','activity_account.xml','AccountActivity','Hồ sơ riêng với họ tên, giới thiệu và thông tin đăng nhập; thay cho hộp thoại chỉ hiển thị UID.','Dữ liệu hồ sơ tham chiếu document users/{uid}. Chỉnh thông tin rồi lưu; chờ xác nhận trước khi đối chiếu các trường displayName và bio trên Firebase Console.'),
 ('Giao diện sao lưu minh họa','BackupActivity.png','activity_backup.xml','BackupActivity','Hai công tắc và nút Sao lưu ngay được giữ để trình bày bố cục. Trạng thái ghi rõ đây là màn minh họa.','Nhấn Sao lưu ngay chỉ mở thông báo giải thích; không tải file lên Firebase Storage, không có sao lưu tự động trong bản hiện tại.'),
 ('Thông báo hoạt động','NotificationsActivity.png','activity_notifications.xml và item_notification.xml','NotificationsActivity','Danh sách các hoạt động lưu trong ứng dụng hoặc Firestore và nút đánh dấu tất cả đã đọc.','Đây là lịch sử hoạt động, không phải thông báo đẩy. Nếu danh sách còn câu sao lưu lên Storage từ dữ liệu cũ, câu đó không chứng minh bản hiện tại đã tải file lên cloud.'),
 ('Thùng rác','TrashActivity.png','activity_trash.xml và item_media_row.xml','TrashActivity và GalleryActivity','Danh sách mục đã chuyển vào thùng rác cùng hành động dọn. Nếu chưa xóa mục nào, màn hình hiển thị trạng thái trống.','Chạm mục để khôi phục hoặc xác nhận xóa vĩnh viễn. Khi mở thùng rác, ứng dụng có xử lý các mục quá 30 ngày; không phải dịch vụ dọn trên server khi app đóng.'),
]
for i,(title,name,layout,activity,meaning,steps) in enumerate(screens,1):
    doc.add_page_break();heading(f'3 {i} '+title,2)
    image_path=SHOTS/name
    if not image_path.exists():raise FileNotFoundError(image_path)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
    run=p.add_run();shape=run.add_picture(str(image_path),width=Cm(8))
    shape._inline.docPr.set('descr',title+' Ảnh chụp ứng dụng chạy trên máy ảo')
    cap=para(f'Hình {i} '+title,True);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para('Bố cục: '+layout+'. Lớp xử lý: '+activity+'.')
    para('Chú thích giao diện: '+meaning)
    para('Thao tác trình bày: '+steps)

# Apply the requested font to every authored run, including headers and captions.
for p in list(doc.paragraphs)+list(sec.footer.paragraphs):
    for run in p.runs:
        run.font.name='Times New Roman';run.font.size=Pt(13);run.font.color.rgb=RGBColor(0,0,0)
doc.save(DOC)
(OUT/'report-source-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(DOC),'java_files':len([p for p in manifest if p['path'].endswith('.java')]),'xml_files':len([p for p in manifest if p['path'].endswith('.xml')]),'figures':len(screens),'source_characters':sum(p['characters'] for p in manifest)},ensure_ascii=False))
